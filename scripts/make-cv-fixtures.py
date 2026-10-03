#!/usr/bin/env python3
"""
Makes the CV files used by the tests (tests/cvs/): realistic CVs of international students, in the formats people really upload
(PDF, Word, plain text) and the ones that go wrong (a scanned page, an old .doc). Run once; the files are committed.
  python3 scripts/make-cv-fixtures.py
"""
import os
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
from docx import Document
from PIL import Image, ImageDraw, ImageFont

OUT = 'tests/cvs/'
os.makedirs(OUT, exist_ok=True)
CVS = {
 'vietnam-finance': dict(name='Nguyen Thi Lan', line='Rotterdam | lan.nguyen@example.com', summary='Financial analyst with two and a half years at a large Vietnamese bank. IFRS reporting, Excel modelling and budgeting. MSc Finance, Erasmus University Rotterdam.',
   jobs=[('Financial Analyst, Vietcombank, Hanoi, Vietnam','Mar 2021 - Aug 2023',['Prepared monthly IFRS reports for the treasury department','Built Excel models for budgeting and forecasting','Reconciled accounts and supported the annual audit'])],
   edu=['MSc Finance, Erasmus University Rotterdam, 2023 - 2025 (cum laude)','BSc Banking, Banking Academy of Vietnam, 2017 - 2021'], skills=['Excel','IFRS','Financial modelling','Budgeting','Power BI'], extra=['Winner, national finance case competition, Vietnam, 2020'], fmt='pdf'),
 'india-btech': dict(name='Arjun Mehta', line='Eindhoven | arjun.mehta@example.com', summary='Software engineer, B.Tech Computer Science. Java, Python, React, SQL. One internship at Infosys.',
   jobs=[('Software Engineering Intern, Infosys, Bengaluru, India','Jan 2024 - Jun 2024',['Built REST APIs in Java and Spring Boot','Wrote SQL queries and dashboards for a retail client','Worked in a scrum team of six'])],
   edu=['B.Tech Computer Science and Engineering, VIT University, 2020 - 2024, CGPA 8.7/10'], skills=['Java','Python','React','SQL','Git','Docker'], extra=[], fmt='docx'),
 'nigeria-accountant': dict(name='Chidinma Okafor', line='Amsterdam | chidinma.okafor@example.com', summary='ACCA accountant with two years in audit at PwC Nigeria. Audit, IFRS, tax and reconciliation.',
   jobs=[('Audit Associate, PwC Nigeria, Lagos, Nigeria','Sep 2021 - Aug 2023',['Performed substantive testing for banking and manufacturing clients','Prepared IFRS disclosures and working papers','Reviewed tax computations'])],
   edu=['BSc Accounting, University of Lagos, 2016 - 2020 (First Class)','ACCA qualification, 2021 - 2023'], skills=['Audit','IFRS','Tax','Excel','SAP'], extra=['Dean\'s list, University of Lagos, 2018 and 2019'], fmt='pdf'),
 'brazil-pt': dict(name='Mariana Souza', line='Utrecht | mariana.souza@example.com', summary='Profissional de marketing digital. Mestrado em Administração. Estagiária na Natura em marketing digital e CRM.',
   jobs=[('Estagiária de Marketing Digital, Natura, São Paulo, Brasil','Jan 2023 - Dez 2023',['Gestão de campanhas de e-mail e CRM','Análise de métricas no Google Analytics','Produção de conteúdo para redes sociais'])],
   edu=['Mestrado em Administração, USP, 2022 - 2024','Bacharelado em Administração, USP, 2018 - 2022'], skills=['CRM','Google Analytics','Excel','Marketing digital'], extra=[], fmt='docx'),
 'china-mech': dict(name='Wei Zhang', line='Delft | wei.zhang@example.com', summary='Mechanical engineer. Master of Engineering. CAD, SolidWorks and finite element analysis. Internship at BYD.',
   jobs=[('Mechanical Engineering Intern, BYD, Shenzhen, China','Jun 2023 - Dec 2023',['Designed battery housing parts in SolidWorks','Ran finite element analysis for thermal loads','Produced technical drawings for production'])],
   edu=['Master of Engineering, Mechanical Engineering, Tsinghua University, 2022 - 2025','Bachelor of Engineering, Harbin Institute of Technology, 2018 - 2022'], skills=['SolidWorks','CAD','ANSYS','MATLAB'], extra=['National scholarship, China, 2021'], fmt='pdf'),
 'elite-heavy': dict(name='Sofia Rossi', line='Amsterdam | sofia.rossi@example.com', summary='Finance graduate with investment banking and consulting experience. Summa cum laude. International case competition winner.',
   jobs=[('Summer Analyst, Goldman Sachs, London, United Kingdom','Jun 2023 - Aug 2023',['Built valuation models for two M&A deals','Prepared pitch materials for senior bankers']),('Intern, McKinsey and Company, Milan, Italy','Jun 2022 - Aug 2022',['Supported a cost transformation project for a bank'])],
   edu=['MSc Finance, London Business School, 2023 - 2025, distinction','BSc Economics, Bocconi University, 2019 - 2022, summa cum laude, GPA 3.9/4.0'], skills=['Financial modelling','Excel','Valuation','PowerPoint'], extra=['First place, Global Investment Banking Case Competition, international, 2022','Chevening Scholarship, 2023'], fmt='pdf'),
 'weak-cv': dict(name='J. Smith', line='', summary='I am hard working and a team player. Looking for any job.',
   jobs=[('Barista, A cafe','2022 - 2023',['Made coffee','Cleaned tables'])], edu=['High school diploma'], skills=[], extra=[], fmt='txt'),
 'table-resume': dict(name='Priya Nair', line='Rotterdam | priya.nair@example.com', summary='Junior data scientist. Python, SQL, machine learning. MSc Data Science.',
   jobs=[('Data Analyst Intern, Flipkart, Bengaluru, India','Jun 2023 - Dec 2023',['Built churn models in Python','Wrote SQL for weekly reporting dashboards'])],
   edu=['MSc Data Science, Eindhoven University of Technology, 2024 - 2026','B.E. Computer Science, PES University, 2019 - 2023'], skills=['Python','SQL','Machine learning','Tableau'], extra=[], fmt='docx-table'),
}

def text_of(c):
    L=[c['name']]; L+= [c['line']] if c['line'] else []
    L+=['', 'Summary', c['summary'], '', 'Experience']
    for t,d,b in c['jobs']: L+=[t, d]+['- '+x for x in b]+['']
    L+=['Education']+c['edu']+['']
    if c['skills']: L+=['Skills', ', '.join(c['skills']), '']
    if c['extra']: L+=['Awards']+c['extra']
    return '\n'.join(L)

def pdf(c, path):
    st=getSampleStyleSheet(); h=ParagraphStyle('h',parent=st['Heading2'],spaceBefore=8)
    doc=SimpleDocTemplate(path,pagesize=A4); S=[Paragraph(c['name'],st['Title'])]
    if c['line']: S.append(Paragraph(c['line'],st['Normal']))
    S+=[Paragraph('Summary',h),Paragraph(c['summary'],st['Normal']),Paragraph('Experience',h)]
    for t,d,b in c['jobs']:
        S+=[Paragraph('<b>%s</b>'%t,st['Normal']),Paragraph(d,st['Normal']),ListFlowable([ListItem(Paragraph(x,st['Normal'])) for x in b],bulletType='bullet'),Spacer(1,6)]
    S.append(Paragraph('Education',h)); S+=[Paragraph(x,st['Normal']) for x in c['edu']]
    if c['skills']: S+=[Paragraph('Skills',h),Paragraph(', '.join(c['skills']),st['Normal'])]
    if c['extra']: S.append(Paragraph('Awards',h)); S+=[Paragraph(x,st['Normal']) for x in c['extra']]
    doc.build(S)

def docx(c, path, table=False):
    d=Document(); d.add_heading(c['name'],0)
    if c['line']: d.add_paragraph(c['line'])
    d.add_heading('Summary',2); d.add_paragraph(c['summary']); d.add_heading('Experience',2)
    for t,dt,b in c['jobs']:
        if table:
            tb=d.add_table(rows=1,cols=2); tb.rows[0].cells[0].text=t; tb.rows[0].cells[1].text=dt
        else:
            p=d.add_paragraph(); p.add_run(t).bold=True; d.add_paragraph(dt)
        for x in b: d.add_paragraph(x,style='List Bullet')
    d.add_heading('Education',2)
    for x in c['edu']: d.add_paragraph(x)
    if c['skills']:
        d.add_heading('Skills',2)
        if table:
            tb=d.add_table(rows=1,cols=len(c['skills']))
            for i,s in enumerate(c['skills']): tb.rows[0].cells[i].text=s
        else: d.add_paragraph(', '.join(c['skills']))
    if c['extra']:
        d.add_heading('Awards',2)
        for x in c['extra']: d.add_paragraph(x)
    d.save(path)

for k,c in CVS.items():
    f=c['fmt']
    if f=='pdf': pdf(c,OUT+k+'.pdf')
    elif f=='docx': docx(c,OUT+k+'.docx')
    elif f=='docx-table': docx(c,OUT+k+'.docx',table=True)
    else: open(OUT+k+'.txt','w').write(text_of(c))
# a scan: the CV as a picture, no text layer
img=Image.new('RGB',(1240,1754),'white'); dr=ImageDraw.Draw(img)
try: font=ImageFont.truetype('/System/Library/Fonts/Helvetica.ttc',34)
except Exception: font=ImageFont.load_default()
y=80
for line in text_of(CVS['vietnam-finance']).split('\n'): dr.text((80,y),line,fill='black',font=font); y+=44
img.save(OUT+'scanned.pdf','PDF',resolution=150)
open(OUT+'old-format.doc','wb').write(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'+b'\x00'*512)
print(sorted(os.listdir(OUT)))
